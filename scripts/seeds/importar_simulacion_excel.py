"""Carga la simulación de `orientacion_vocacional_simulacion.xlsx` como datos de prueba.

Cada fila del Excel es una aplicación de la prueba de Competencias hecha con el banco
anterior de 30 ítems (5 por campo). Por cada aplicación se crea un estudiante de
simulación, la aplicación completada y sus respuestas, y se calcula el perfil con el
motor actual.

El cálculo se hace con sólo esos 30 ítems activos: con el banco actual de 64, el máximo
posible por campo incluiría preguntas que la persona nunca vio y los porcentajes
saldrían falsamente bajos (y más bajos en los campos con más ítems). Al terminar, el
banco completo se vuelve a activar.

Es idempotente (omite estudiantes que ya existan) y no corre en producción.

Uso:
    python scripts/seeds/importar_simulacion_excel.py
    python scripts/seeds/importar_simulacion_excel.py --archivo ruta.xlsx --dry-run
    python scripts/seeds/importar_simulacion_excel.py --recalcular   # rehace los perfiles ya cargados
"""

from __future__ import annotations

import argparse
import os
import secrets
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(RAIZ / "backend"))

from app import create_app  # noqa: E402
from app.extensions import db  # noqa: E402
from app.models.aplicacion import Aplicacion, ConfiguracionAplicacion, Respuesta  # noqa: E402
from app.models.instrumento import Item  # noqa: E402
from app.models.usuario import Rol, Usuario  # noqa: E402
from app.services.auth_service import AuthService  # noqa: E402
from app.services.psicometrico_service import PsicometricoService  # noqa: E402

CAMPOS = {"R": "Realista", "I": "Investigador", "A": "Artístico",
          "S": "Social", "E": "Emprendedor", "C": "Convencional"}
DOMINIO = "simulacion.local"


def leer_hoja(libro, nombre: str) -> list[dict]:
    filas = libro[nombre].iter_rows(values_only=True)
    encabezado = next(filas)
    return [dict(zip(encabezado, f)) for f in filas if f and f[0] is not None]


def fecha(valor) -> datetime | None:
    if valor in (None, ""):
        return None
    if isinstance(valor, datetime):
        return valor
    return datetime.fromisoformat(str(valor))


def configuracion_competencias() -> ConfiguracionAplicacion:
    for cfg in ConfiguracionAplicacion.query.filter_by(activa=True).all():
        if "competencia" in (cfg.instrumento.nombre or "").lower():
            return cfg
    raise SystemExit("No hay una configuración activa de la prueba de Competencias.")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--archivo", default=str(RAIZ / "orientacion_vocacional_simulacion.xlsx"))
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--recalcular", action="store_true",
                        help="recalcula los perfiles de las aplicaciones de simulación ya cargadas")
    args = parser.parse_args()

    entorno = os.getenv("FLASK_ENV", "development")
    if entorno == "production":
        raise SystemExit("Este script carga datos de simulación: no se ejecuta en producción.")

    libro = load_workbook(args.archivo, read_only=True, data_only=True)
    identificacion = leer_hoja(libro, "Identificacion")
    aplicacion_xls = {str(f["id_aplicacion"]): f for f in leer_hoja(libro, "Aplicacion")}
    competencias_xls = {str(f["id_aplicacion"]): f for f in leer_hoja(libro, "Competencias")}
    respuestas_xls: dict[str, list[dict]] = {}
    for f in leer_hoja(libro, "Respuestas"):
        respuestas_xls.setdefault(str(f["id_aplicacion"]), []).append(f)

    app = create_app(entorno)
    with app.app_context():
        cfg = configuracion_competencias()
        rol = Rol.query.filter_by(nombre="estudiante").first()
        items = {i.codigo: i for i in Item.query.filter_by(tipo="juicio_situacional").all()}
        codigos_aplicados = {r["codigo_item"] for filas in respuestas_xls.values() for r in filas}
        faltantes = codigos_aplicados - set(items)
        if faltantes:
            raise SystemExit(f"Códigos del Excel que no están en el banco: {sorted(faltantes)}")

        nuevas: list[tuple[Aplicacion, str]] = []
        omitidas = 0
        if args.recalcular:
            # Rehace los perfiles de las aplicaciones de simulación existentes (p. ej. tras
            # cambiar el motor de puntaje), otra vez con el banco de 30 ítems.
            for fila in identificacion:
                u = Usuario.query.filter_by(email=f"sim{fila['id_estudiante']}@{DOMINIO}").first()
                ap = u.aplicaciones.filter_by(estado="completada").first() if u else None
                if ap:
                    nuevas.append((ap, str(fila["id_aplicacion"])))
            identificacion = []
        for fila in identificacion:
            id_app, id_est = str(fila["id_aplicacion"]), str(fila["id_estudiante"])
            email = f"sim{id_est}@{DOMINIO}"
            if Usuario.query.filter_by(email=email).first():
                omitidas += 1
                continue
            meta = aplicacion_xls.get(id_app, {})
            inicio = fecha(meta.get("fecha_inicio")) or fecha(fila.get("fecha_aplicacion"))
            fin = fecha(meta.get("fecha_fin")) or inicio
            usuario = Usuario(
                email=email,
                # Contraseña aleatoria: son cuentas de datos, no para iniciar sesión.
                password_hash=AuthService.hash_password(secrets.token_urlsafe(24)),
                nombres="Estudiante", apellidos=f"Simulación {id_est}",
                rol_id=rol.id, cohorte="SIM-2026", activo=True, created_at=inicio,
            )
            db.session.add(usuario)
            db.session.flush()
            aplicacion = Aplicacion(
                usuario_id=usuario.id, configuracion_id=cfg.id, estado="completada",
                progreso=100, fecha_inicio=inicio, fecha_fin=fin, created_at=inicio,
                user_agent=meta.get("dispositivo"),
            )
            db.session.add(aplicacion)
            db.session.flush()
            for r in respuestas_xls.get(id_app, []):
                db.session.add(Respuesta(
                    aplicacion_id=aplicacion.id,
                    item_id=items[r["codigo_item"]].id,
                    valor=int(r["respuesta_original"]),
                ))
            nuevas.append((aplicacion, id_app))

        print(f"Aplicaciones a procesar: {len(nuevas)} | ya existían (omitidas): {omitidas}")
        if args.dry_run:
            db.session.rollback()
            print("Dry-run: no se modificó la base de datos.")
            return
        db.session.commit()

        # Calcular con el banco que se aplicó (30 ítems) y luego restaurar el banco completo.
        apagados = [i for c, i in items.items() if i.activo and c not in codigos_aplicados]
        try:
            for i in apagados:
                i.activo = False
            db.session.commit()
            for aplicacion, _ in nuevas:
                PsicometricoService.calcular_puntajes_dimension(aplicacion.id)
                PsicometricoService.generar_perfil(aplicacion.id)
        finally:
            for i in apagados:
                i.activo = True
            db.session.commit()

        # Verificación: el % de competencias calculado debe coincidir con el del Excel.
        coinciden = 0
        for aplicacion, id_app in nuevas:
            areas = (aplicacion.perfil.datos_json or {}).get("puntajes_por_area", {})
            esperado = competencias_xls.get(id_app, {})
            if all(round(areas.get(nombre, {}).get("competencia_porcentaje", -1)) == int(esperado.get(f"porcentaje_{letra}", -2))
                   for letra, nombre in CAMPOS.items()):
                coinciden += 1
        print(f"Perfiles calculados: {len(nuevas)} | porcentajes iguales al Excel: {coinciden}/{len(nuevas)}")


if __name__ == "__main__":
    main()
